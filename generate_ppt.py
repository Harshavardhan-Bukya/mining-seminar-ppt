from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE
import os

# ============================================================
# MINING SEMINAR PPT GENERATOR
# Civil Engineering - 15 Slides
# ============================================================

OUTPUT = "Mining_Seminar_15_Slides.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ---------------- COLORS ----------------

NAVY = RGBColor(25, 39, 47)
DARK = RGBColor(48, 55, 58)
GOLD = RGBColor(218, 158, 43)
SAND = RGBColor(239, 233, 218)
LIGHT = RGBColor(242, 244, 242)
WHITE = RGBColor(255, 255, 255)
GREY = RGBColor(105, 112, 114)
GREEN = RGBColor(71, 117, 82)
BLUE = RGBColor(69, 104, 124)
RED = RGBColor(171, 75, 61)
BROWN = RGBColor(115, 91, 61)

# ---------------- BASIC FUNCTIONS ----------------

def set_background(slide, color=WHITE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0, 0,
        prs.slide_width,
        prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()

    slide.shapes._spTree.remove(shape._element)
    slide.shapes._spTree.insert(2, shape._element)


def add_text(slide, text, x, y, w, h,
             size=18, color=DARK,
             bold=False, align=PP_ALIGN.LEFT,
             font="Aptos"):

    box = slide.shapes.add_textbox(
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.alignment = align

    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color

    return box


def add_header(slide, heading, number):
    add_text(
        slide,
        heading,
        .65, .30, 10.8, .55,
        size=27,
        color=NAVY,
        bold=True,
        font="Aptos Display"
    )

    add_text(
        slide,
        f"{number:02d}",
        12.0, .32, .65, .35,
        size=14,
        color=GOLD,
        bold=True,
        align=PP_ALIGN.RIGHT
    )

    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(.65), Inches(1.0),
        Inches(12.0), Inches(.035)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD
    line.line.fill.background()


def add_footer(slide):
    add_text(
        slide,
        "CIVIL ENGINEERING SEMINAR  •  MINING",
        .65, 7.08, 12, .2,
        size=8,
        color=GREY
    )


def add_bullets(slide, items, x, y, w, h, size=16):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    tf = box.text_frame
    tf.word_wrap = True

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "• " + item
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = DARK
        p.space_after = Pt(10)

    return box


def add_card(slide, x, y, w, h, heading, body, accent=GOLD):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = RGBColor(215, 220, 218)

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y),
        Inches(.08), Inches(h)
    )

    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    add_text(
        slide, heading,
        x + .25, y + .18,
        w - .45, .4,
        size=16,
        color=NAVY,
        bold=True
    )

    add_text(
        slide, body,
        x + .25, y + .68,
        w - .45, h - .8,
        size=12.5,
        color=DARK
    )


# ============================================================
# VISUALS / DIAGRAMS
# ============================================================

def mining_scene(slide, x, y, w, h):
    """Stylized open-cast mine illustration."""

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(228, 235, 236)
    frame.line.color.rgb = RGBColor(205, 213, 214)

    # Ground
    ground = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y + h*.43),
        Inches(w),
        Inches(h*.57)
    )

    ground.fill.solid()
    ground.fill.fore_color.rgb = BROWN
    ground.line.fill.background()

    # Pit benches
    for i, fraction in enumerate([.88, .70, .52, .34]):
        pit = slide.shapes.add_shape(
            MSO_SHAPE.ISOSCELES_TRIANGLE,
            Inches(x + w*(1-fraction)/2),
            Inches(y + h*(.51 + i*.105)),
            Inches(w*fraction),
            Inches(h*.22)
        )

        pit.rotation = 180
        pit.fill.solid()
        pit.fill.fore_color.rgb = RGBColor(
            83 + i*8,
            70 + i*7,
            55 + i*5
        )
        pit.line.fill.background()

    # Haul road
    road = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM,
        Inches(x+w*.12),
        Inches(y+h*.66),
        Inches(w*.76),
        Inches(.22)
    )

    road.fill.solid()
    road.fill.fore_color.rgb = RGBColor(184, 171, 143)
    road.line.fill.background()

    # Equipment
    for dx, dy in [(.22,.62), (.54,.73), (.72,.57)]:
        truck = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x+w*dx),
            Inches(y+h*dy),
            Inches(.38),
            Inches(.14)
        )

        truck.fill.solid()
        truck.fill.fore_color.rgb = GOLD
        truck.line.fill.background()

    add_text(
        slide,
        "OPEN-CAST MINE — ILLUSTRATIVE ENGINEERING VISUAL",
        x+.2, y+h-.42,
        w-.4, .25,
        size=9,
        color=NAVY,
        align=PP_ALIGN.CENTER
    )


def underground_scene(slide, x, y, w, h):
    """Stylized underground mine diagram."""

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(239, 234, 220)
    frame.line.color.rgb = RGBColor(210, 201, 180)

    # Surface
    surface = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x),
        Inches(y+h*.20),
        Inches(w),
        Inches(.18)
    )

    surface.fill.solid()
    surface.fill.fore_color.rgb = GREEN
    surface.line.fill.background()

    # Shaft
    shaft = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(x+w*.46),
        Inches(y+h*.28),
        Inches(.30),
        Inches(h*.50)
    )

    shaft.fill.solid()
    shaft.fill.fore_color.rgb = NAVY
    shaft.line.fill.background()

    # Tunnels
    for yy in [.51, .66, .81]:
        tunnel = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x+w*.16),
            Inches(y+h*yy),
            Inches(w*.68),
            Inches(.10)
        )

        tunnel.fill.solid()
        tunnel.fill.fore_color.rgb = NAVY
        tunnel.line.fill.background()

    # Ore body
    ore = slide.shapes.add_shape(
        MSO_SHAPE.DIAMOND,
        Inches(x+w*.62),
        Inches(y+h*.47),
        Inches(.52),
        Inches(.72)
    )

    ore.fill.solid()
    ore.fill.fore_color.rgb = GOLD
    ore.line.fill.background()

    add_text(
        slide,
        "UNDERGROUND MINE — ILLUSTRATIVE LAYOUT",
        x+.2, y+h-.42,
        w-.4, .25,
        size=9,
        color=NAVY,
        align=PP_ALIGN.CENTER
    )


def equipment_visual(slide, x, y, w, h):
    """Stylized mining machinery visual."""

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(230, 234, 232)
    frame.line.color.rgb = RGBColor(205, 212, 210)

    labels = [
        ("EXCAVATOR", GOLD),
        ("HAUL TRUCK", BLUE),
        ("DRILL RIG", RED)
    ]

    for i, (label, accent) in enumerate(labels):

        bx = x + .45 + i*w*.29
        by = y + h*.52

        body = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(bx),
            Inches(by),
            Inches(w*.18),
            Inches(h*.08)
        )

        body.fill.solid()
        body.fill.fore_color.rgb = accent
        body.line.fill.background()

        # Wheels
        for wx in [bx+.05, bx+w*.12]:
            wheel = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(wx),
                Inches(by+h*.11),
                Inches(.25),
                Inches(.25)
            )

            wheel.fill.solid()
            wheel.fill.fore_color.rgb = NAVY
            wheel.line.fill.background()

        add_text(
            slide,
            label,
            bx-.15,
            y+h*.78,
            w*.30,
            .35,
            size=9,
            color=NAVY,
            bold=True,
            align=PP_ALIGN.CENTER
        )


def slope_visual(slide, x, y, w, h):
    """Stylized open-pit bench and slope diagram."""

    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y),
        Inches(w), Inches(h)
    )

    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(241, 239, 232)
    frame.line.color.rgb = RGBColor(215, 210, 198)

    for i in range(5):
        bench = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x+.5+i*.55),
            Inches(y+h*.70-i*.31),
            Inches(w*.65-i*.07),
            Inches(.12)
        )

        bench.fill.solid()
        bench.fill.fore_color.rgb = GOLD
        bench.line.fill.background()

    add_text(
        slide,
        "BENCHES + SLOPE FACE + CATCH BERM",
        x+.25, y+h-.42,
        w-.5, .25,
        size=9,
        color=NAVY,
        align=PP_ALIGN.CENTER
    )


def process_flow(slide, labels, x, y, w):
    gap = .12
    bw = (w - gap*(len(labels)-1)) / len(labels)

    for i, label in enumerate(labels):

        bx = x + i*(bw+gap)

        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(bx), Inches(y),
            Inches(bw), Inches(.72)
        )

        box.fill.solid()
        box.fill.fore_color.rgb = NAVY
        box.line.fill.background()

        add_text(
            slide,
            label,
            bx+.04, y+.12,
            bw-.08, .45,
            size=9.5,
            color=WHITE,
            bold=True,
            align=PP_ALIGN.CENTER
        )

        if i < len(labels)-1:

            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(bx+bw+.01),
                Inches(y+.22),
                Inches(.10),
                Inches(.25)
            )

            arrow.fill.solid()
            arrow.fill.fore_color.rgb = GOLD
            arrow.line.fill.background()


# ============================================================
# SLIDE 1 — TITLE
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, NAVY)

mining_scene(slide, 7.15, 1.0, 5.45, 5.55)

add_text(
    slide,
    "MINING",
    .8, 1.2, 5.8, 1,
    size=52,
    color=WHITE,
    bold=True,
    font="Aptos Display"
)

add_text(
    slide,
    "Seminar Presentation",
    .82, 2.5, 5.7, .6,
    size=24,
    color=GOLD
)

add_text(
    slide,
    "Presented by: [Your Name]\n"
    "Roll No.: [Your Roll Number]\n"
    "Department of Civil Engineering\n"
    "[College Name]",
    .82, 4.0, 5.7, 1.6,
    size=15,
    color=WHITE
)


# ============================================================
# SLIDE 2 — INTRODUCTION
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Introduction to Mining", 2)

add_bullets(slide, [
    "Mining is the extraction of useful minerals and geological materials from the Earth.",
    "It supplies raw materials for construction, manufacturing, energy and infrastructure.",
    "The mining method depends on deposit depth, geometry, rock conditions and economics.",
    "Mining combines geology, excavation, transportation, safety and environmental management."
], .8, 1.35, 6.0, 4.9, 16)

mining_scene(slide, 7.25, 1.4, 5.25, 4.65)

add_footer(slide)


# ============================================================
# SLIDE 3 — OBJECTIVES AND IMPORTANCE
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Objectives and Importance of Mining", 3)

add_card(
    slide, .8, 1.4, 5.55, 1.8,
    "Resource Extraction",
    "Recover useful mineral resources safely and economically.",
    GOLD
)

add_card(
    slide, 6.75, 1.4, 5.55, 1.8,
    "Infrastructure",
    "Provide materials such as iron ore, limestone and aggregates.",
    BLUE
)

add_card(
    slide, .8, 3.45, 5.55, 1.8,
    "Industrial Development",
    "Support steel, cement, power and manufacturing industries.",
    GREEN
)

add_card(
    slide, 6.75, 3.45, 5.55, 1.8,
    "Employment & Economy",
    "Generate direct and indirect employment and economic activity.",
    RED
)

process_flow(
    slide,
    ["MINERALS", "INDUSTRY", "INFRASTRUCTURE"],
    2.25, 6.05, 8.9
)

add_footer(slide)


# ============================================================
# SLIDE 4 — TYPES OF MINING
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Types of Mining", 4)

mining_scene(
    slide, .8, 1.35, 5.45, 4.95
)

underground_scene(
    slide, 7.05, 1.35, 5.45, 4.95
)

add_text(
    slide,
    "SURFACE / OPEN-CAST",
    1.0, 6.35, 5.0, .3,
    size=14, color=NAVY, bold=True,
    align=PP_ALIGN.CENTER
)

add_text(
    slide,
    "UNDERGROUND MINING",
    7.25, 6.35, 5.0, .3,
    size=14, color=NAVY, bold=True,
    align=PP_ALIGN.CENTER
)

add_footer(slide)


# ============================================================
# SLIDE 5 — MINING METHODS
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Mining Methods", 5)

add_card(
    slide, .8, 1.4, 5.55, 1.75,
    "Open-Pit Mining",
    "Large surface excavation using benches and haul roads.",
    GOLD
)

add_card(
    slide, 6.75, 1.4, 5.55, 1.75,
    "Strip Mining",
    "Overburden is removed in strips; suitable for relatively flat deposits.",
    BLUE
)

add_card(
    slide, .8, 3.45, 5.55, 1.75,
    "Room-and-Pillar",
    "Rooms are extracted while pillars provide roof support.",
    GREEN
)

add_card(
    slide, 6.75, 3.45, 5.55, 1.75,
    "Longwall Mining",
    "Mechanized extraction across a long panel, commonly used in coal mining.",
    RED
)

add_footer(slide)


# ============================================================
# SLIDE 6 — MINERALS
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Major Minerals and Resources", 6)

minerals = [
    ("COAL", "Power and industrial fuel", DARK),
    ("IRON ORE", "Steel production", RED),
    ("LIMESTONE", "Cement and construction", GOLD),
    ("BAUXITE", "Aluminium production", BLUE),
    ("COPPER", "Electrical and infrastructure", GREEN),
    ("GOLD", "Jewellery and technology", RGBColor(150,120,45))
]

for i, (name, use, color) in enumerate(minerals):

    x = .75 + (i % 3) * 4.15
    y = 1.45 + (i // 3) * 2.35

    circle = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x+.1), Inches(y+.1),
        Inches(.72), Inches(.72)
    )

    circle.fill.solid()
    circle.fill.fore_color.rgb = color
    circle.line.fill.background()

    add_text(
        slide, name,
        x+1.0, y+.03,
        2.8, .42,
        size=15, color=NAVY, bold=True
    )

    add_text(
        slide, use,
        x+1.0, y+.52,
        2.85, .65,
        size=11, color=GREY
    )

add_card(
    slide,
    1.15, 6.1, 11.0, .62,
    "Civil Engineering Connection",
    "Mineral resources support cement, steel, roads, buildings, bridges and other infrastructure.",
    GOLD
)

add_footer(slide)


# ============================================================
# SLIDE 7 — MINING OPERATIONS
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Mining Operations", 7)

process_flow(
    slide,
    [
        "EXPLORATION",
        "PLANNING",
        "DRILLING",
        "BLASTING",
        "EXCAVATION",
        "HAULING",
        "PROCESSING"
    ],
    .75, 1.45, 11.85
)

add_card(
    slide, .8, 3.0, 3.65, 1.55,
    "Exploration",
    "Locate and characterize the mineral deposit.",
    GOLD
)

add_card(
    slide, 4.85, 3.0, 3.65, 1.55,
    "Extraction",
    "Drill, blast, excavate and load material.",
    RED
)

add_card(
    slide, 8.9, 3.0, 3.65, 1.55,
    "Transport & Processing",
    "Move material and prepare the product for use.",
    BLUE
)

equipment_visual(
    slide, 3.45, 5.05, 6.4, 1.55
)

add_footer(slide)


# ============================================================
# SLIDE 8 — EQUIPMENT
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Major Mining Equipment", 8)

equipment_visual(
    slide, .8, 1.35, 6.15, 4.95
)

add_bullets(slide, [
    "Excavators: digging and loading.",
    "Haul trucks: transporting waste and ore.",
    "Drilling rigs: creating blast holes.",
    "Bulldozers and graders: earthwork and road maintenance.",
    "Crushers and conveyors: size reduction and material transport."
], 7.35, 1.5, 5.0, 4.6, 15)

add_footer(slide)


# ============================================================
# SLIDE 9 — MINE PLANNING
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Mine Planning and Design", 9)

add_bullets(slide, [
    "Geological and geotechnical investigation defines the deposit and ground conditions.",
    "Mine geometry includes benches, pit limits or underground openings.",
    "Haul roads require suitable alignment, gradients, width and drainage.",
    "Production planning coordinates equipment, material movement and processing.",
    "Monitoring and design review help manage changing ground conditions."
], .8, 1.4, 6.1, 4.8, 15)

mining_scene(
    slide, 7.25, 1.45, 5.25, 4.75
)

add_footer(slide)


# ============================================================
# SLIDE 10 — SLOPE STABILITY & SAFETY
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Slope Stability and Mining Safety", 10)

slope_visual(
    slide, .8, 1.35, 6.0, 4.95
)

add_bullets(slide, [
    "Assess rock mass properties, groundwater and geological discontinuities.",
    "Use suitable bench geometry, berms and catchment areas.",
    "Control blasting to limit damage and vibration.",
    "Use PPE, traffic controls, equipment inspections and safe procedures.",
    "Maintain emergency response plans and regular safety training."
], 7.2, 1.5, 5.1, 4.7, 15)

add_footer(slide)


# ============================================================
# SLIDE 11 — ENVIRONMENT
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Environmental Impacts of Mining", 11)

add_card(
    slide, .8, 1.4, 5.55, 1.75,
    "Land Disturbance",
    "Excavation, waste dumps and infrastructure modify the landscape.",
    RED
)

add_card(
    slide, 6.75, 1.4, 5.55, 1.75,
    "Air Pollution",
    "Dust and equipment emissions can affect air quality.",
    GOLD
)

add_card(
    slide, .8, 3.45, 5.55, 1.75,
    "Water Impacts",
    "Sediment, drainage and contaminants may affect water resources.",
    BLUE
)

add_card(
    slide, 6.75, 3.45, 5.55, 1.75,
    "Noise & Vibration",
    "Machinery and blasting generate noise and ground vibration.",
    DARK
)

process_flow(
    slide,
    ["ASSESS", "CONTROL", "MONITOR", "RECLAIM"],
    1.25, 6.05, 10.85
)

add_footer(slide)


# ============================================================
# SLIDE 12 — SUSTAINABILITY
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Sustainable Mining", 12)

process_flow(
    slide,
    ["PLAN", "MINIMIZE", "RECOVER", "REUSE", "RESTORE"],
    .85, 1.45, 11.65
)

add_card(
    slide, .8, 3.0, 5.55, 1.5,
    "Efficient Resource Use",
    "Optimize extraction to reduce unnecessary waste and energy use.",
    GOLD
)

add_card(
    slide, 6.75, 3.0, 5.55, 1.5,
    "Water & Waste Management",
    "Recycle water where feasible and manage mine waste responsibly.",
    BLUE
)

add_card(
    slide, .8, 4.8, 5.55, 1.5,
    "Pollution Control",
    "Apply dust suppression, noise control and emission management.",
    GREEN
)

add_card(
    slide, 6.75, 4.8, 5.55, 1.5,
    "Progressive Reclamation",
    "Restore suitable areas during the mine life rather than waiting until closure.",
    RED
)

add_footer(slide)


# ============================================================
# SLIDE 13 — RECLAMATION
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Mine Reclamation", 13)

mining_scene(
    slide, .8, 1.4, 5.45, 4.8
)

slope_visual(
    slide, 7.05, 1.4, 5.45, 4.8
)

add_text(
    slide,
    "DISTURBED LAND",
    1.0, 6.30, 5.0, .3,
    size=14, color=RED, bold=True,
    align=PP_ALIGN.CENTER
)

add_text(
    slide,
    "RESTORED LAND",
    7.25, 6.30, 5.0, .3,
    size=14, color=GREEN, bold=True,
    align=PP_ALIGN.CENTER
)

process_flow(
    slide,
    ["BACKFILL", "RESHAPE", "TOPSOIL", "REVEGETATE", "MONITOR"],
    .95, 6.65, 11.35
)

add_footer(slide)


# ============================================================
# SLIDE 14 — MODERN TECHNOLOGY
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide)
add_header(slide, "Modern Mining Technology and Future Scope", 14)

add_card(
    slide, .75, 1.35, 3.75, 1.65,
    "GPS & GIS",
    "Precise positioning, mapping and spatial mine planning.",
    BLUE
)

add_card(
    slide, 4.8, 1.35, 3.75, 1.65,
    "Drones",
    "Rapid surveying, inspection and terrain mapping.",
    GOLD
)

add_card(
    slide, 8.85, 1.35, 3.75, 1.65,
    "IoT Sensors",
    "Real-time monitoring of equipment and mine conditions.",
    GREEN
)

add_card(
    slide, .75, 3.35, 3.75, 1.65,
    "AI & Analytics",
    "Predictive maintenance and data-driven decisions.",
    RED
)

add_card(
    slide, 4.8, 3.35, 3.75, 1.65,
    "Automation",
    "Remote and autonomous equipment can improve safety and productivity.",
    DARK
)

add_card(
    slide, 8.85, 3.35, 3.75, 1.65,
    "Future Mining",
    "Digital, connected and lower-impact mining operations.",
    GOLD
)

equipment_visual(
    slide, 2.0, 5.35, 9.3, 1.05
)

add_footer(slide)


# ============================================================
# SLIDE 15 — CONCLUSION & REFERENCES
# ============================================================

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_background(slide, NAVY)

add_text(
    slide,
    "CONCLUSION & REFERENCES",
    .8, .6, 11.7, .6,
    size=30,
    color=WHITE,
    bold=True,
    align=PP_ALIGN.CENTER
)

# Conclusion panel
panel = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(1.0), Inches(1.55),
    Inches(5.35), Inches(3.5)
)

panel.fill.solid()
panel.fill.fore_color.rgb = RGBColor(239, 242, 240)
panel.line.color.rgb = GOLD

add_text(
    slide,
    "CONCLUSION",
    1.25, 1.8, 4.8, .45,
    size=18,
    color=NAVY,
    bold=True
)

add_text(
    slide,
    "Mining provides essential resources for infrastructure and industry.\n\n"
    "Safe mine planning, geotechnical control and responsible operations are fundamental.\n\n"
    "Sustainable practices and modern technology can improve safety, efficiency and environmental performance.",
    1.25, 2.4, 4.8, 2.3,
    size=12.5,
    color=DARK
)

# References panel
panel = slide.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(6.95), Inches(1.55),
    Inches(5.35), Inches(3.5)
)

panel.fill.solid()
panel.fill.fore_color.rgb = RGBColor(239, 242, 240)
panel.line.color.rgb = BLUE

add_text(
    slide,
    "REFERENCES",
    7.2, 1.8, 4.8, .45,
    size=18,
    color=NAVY,
    bold=True
)

add_text(
    slide,
    "• Mining engineering textbooks\n"
    "• Government mining and geological agencies\n"
    "• Peer-reviewed research papers\n"
    "• University technical resources\n"
    "• Official industry guidance\n"
    "• GitHub project resources, where applicable",
    7.2, 2.4, 4.8, 2.3,
    size=12.5,
    color=DARK
)

add_text(
    slide,
    "THANK YOU",
    1.0, 5.55, 11.3, .65,
    size=32,
    color=GOLD,
    bold=True,
    align=PP_ALIGN.CENTER
)

add_text(
    slide,
    "Questions & Discussion",
    1.0, 6.25, 11.3, .4,
    size=15,
    color=WHITE,
    align=PP_ALIGN.CENTER
)


# ============================================================
# SAVE PPT
# ============================================================

prs.save(OUTPUT)

print("=" * 60)
print("MINING SEMINAR PPT CREATED SUCCESSFULLY")
print("=" * 60)
print(f"File: {OUTPUT}")
print(f"Slides: {len(prs.slides)}")
print("=" * 60)
